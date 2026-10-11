"""
Calculate Excel formulas that were saved without their results.

Excel and LibreOffice save each formula's result in the file, and Datasquint reads those results. Some
programs (report exporters, Python libraries) save only the formulas; their cells would load empty. This
module finds such formulas and calculates them, trying in turn:

  1. Microsoft Excel (Windows, if installed): opens the file in a hidden Excel, recalculates, saves a copy.
  2. LibreOffice (if installed): the same, headless.
  3. pycel (pure Python, in requirements.txt): slower, so only for small files, and it knows fewer
     functions: up to PYCEL_MAX_FORMULAS formulas (5,000), PYCEL_MAX_MB (20 MB), PYCEL_TIME_LIMIT (120 s).

The result is a copy of the workbook in which those cells hold values; the original upload is untouched.
Generated Python scripts embed this file, so they calculate the same way.
"""
import contextlib
import io
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

log = logging.getLogger("datasquint.recalc")

ENGINES = [e.strip() for e in os.getenv("RECALC_ENGINES", "excel,libreoffice,pycel").lower().split(",") if e.strip()]
PYCEL_MAX_FORMULAS = int(os.getenv("PYCEL_MAX_FORMULAS", 5000))
PYCEL_MAX_MB = float(os.getenv("PYCEL_MAX_MB", 20))
PYCEL_TIME_LIMIT = float(os.getenv("PYCEL_TIME_LIMIT", 120))
NAMES = {"excel": "Microsoft Excel", "libreoffice": "LibreOffice", "pycel": "the built-in calculator"}


def _short(e):
    text = str(e).strip().splitlines()[0] if str(e).strip() else type(e).__name__
    return text[:200]


# ---------------------------------------------------------------- finding formulas without results
_SHEET_RE = re.compile(rb'<sheet\b[^>]*?\bname="([^"]*)"[^>]*?\br:id="([^"]+)"|<sheet\b[^>]*?\br:id="([^"]+)"[^>]*?\bname="([^"]*)"')
_REL_RE = re.compile(rb'<Relationship\b[^>]*?\bId="([^"]+)"[^>]*?\bTarget="([^"]+)"|<Relationship\b[^>]*?\bTarget="([^"]+)"[^>]*?\bId="([^"]+)"')
_CACHED_RE = re.compile(rb"(?:</f>|<f\b[^>]*/>)\s*<v>[^<]")   # a formula followed by a non-empty saved value
_CELL_RE = re.compile(rb'<c\b([^>]*)>(.*?)</c>', re.S)
_REF_RE = re.compile(rb'\br="([A-Z]+[0-9]+)"')
SCAN_BYTES = 16 << 20   # per sheet: missing results show up throughout a formula column, so a sample is enough


def _unescape(b):
    return (b.decode("utf-8", "replace").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
            .replace("&apos;", "'").replace("&amp;", "&"))


def _sheet_members(z):
    """[(sheet name, zip member)] from the workbook's own index."""
    wb, rels = z.read("xl/workbook.xml"), z.read("xl/_rels/workbook.xml.rels")
    targets = {}
    for m in _REL_RE.finditer(rels):
        rid, tgt = (m.group(1), m.group(2)) if m.group(1) else (m.group(4), m.group(3))
        targets[rid.decode()] = tgt.decode().lstrip("/")
    out, names = [], set(z.namelist())
    for m in _SHEET_RE.finditer(wb):
        name, rid = (m.group(1), m.group(2)) if m.group(1) else (m.group(4), m.group(3))
        tgt = targets.get(rid.decode(), "")
        member = tgt if tgt.startswith("xl/") else "xl/" + tgt
        if member in names:
            out.append((_unescape(name), member))
    return out


def uncached(path):
    """{sheet name: formula cells without a saved result} (from a sample of each sheet). Never raises."""
    try:
        with zipfile.ZipFile(path) as z:
            out = {}
            for name, member in _sheet_members(z):
                total = cached = seen = 0
                carry = b""
                with z.open(member) as fh:
                    while True:
                        chunk = fh.read(4 << 20) if seen < SCAN_BYTES else b""
                        seen += len(chunk)
                        buf = carry + chunk
                        cut = len(buf) if not chunk else max(0, len(buf) - 400)
                        head = buf[:cut]
                        total += head.count(b"<f>") + head.count(b"<f ")
                        cached += sum(1 for x in _CACHED_RE.finditer(buf) if x.start() < cut)
                        carry = buf[cut:]
                        if not chunk:
                            break
                if total > cached:
                    out[name] = total - cached
            return out
    except Exception as e:
        log.info("Couldn't check %s for formulas without results: %s", os.path.basename(path), _short(e))
        return {}


def uncached_cells(path, limit):
    """{sheet: [cell refs]} of all formulas without a saved result, or None if there are more than `limit`."""
    out, n = {}, 0
    with zipfile.ZipFile(path) as z:
        for name, member in _sheet_members(z):
            data = z.read(member)
            for m in _CELL_RE.finditer(data):
                body = m.group(2)
                if (b"<f>" in body or b"<f " in body) and not _CACHED_RE.search(body):
                    ref = _REF_RE.search(m.group(1))
                    if ref:
                        out.setdefault(name, []).append(ref.group(1).decode())
                        n += 1
                        if n > limit:
                            return None
    return out


# ---------------------------------------------------------------- engines
def find_excel():
    """True if Microsoft Excel can be automated here (Windows, Excel installed, pywin32 available)."""
    if os.name != "nt":
        return False
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Excel.Application"))
        import importlib.util
        return importlib.util.find_spec("win32com") is not None   # pywin32
    except Exception:
        return False


def find_libreoffice():
    """Path to LibreOffice's soffice program, or None. LIBREOFFICE in .env overrides the search."""
    if os.getenv("LIBREOFFICE"):
        return os.getenv("LIBREOFFICE") if os.path.exists(os.getenv("LIBREOFFICE")) else None
    for name in ("soffice", "libreoffice"):
        if shutil.which(name):
            return shutil.which(name)
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice", "/opt/libreoffice/program/soffice"):
        if os.path.exists(p):
            return p
    return None


def find_pycel():
    import importlib.util
    return importlib.util.find_spec("pycel") is not None


_EXCEL_WORKER = r"""
import os, sys, pythoncom, win32com.client
src, dst = sys.argv[1], sys.argv[2]
pythoncom.CoInitialize()
xl = win32com.client.DispatchEx("Excel.Application")   # a separate, hidden Excel
try:
    xl.Visible = False; xl.DisplayAlerts = False; xl.ScreenUpdating = False; xl.EnableEvents = False
    xl.AutomationSecurity = 3                                # msoAutomationSecurityForceDisable: no macros
    wb = xl.Workbooks.Open(src, UpdateLinks=0, ReadOnly=True, IgnoreReadOnlyRecommended=True)
    xl.CalculateFull()
    wb.SaveAs(dst, FileFormat=52 if dst.lower().endswith(".xlsm") else 51)
    wb.Close(SaveChanges=False)
finally:
    xl.Quit()
    pythoncom.CoUninitialize()
"""


def with_excel(path, out_dir, timeout=180):
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.abspath(os.path.join(out_dir, os.path.basename(path)))
    r = subprocess.run([sys.executable, "-c", _EXCEL_WORKER, os.path.abspath(path), out],
                       capture_output=True, text=True, timeout=timeout)
    if r.returncode or not os.path.exists(out):
        raise RuntimeError((r.stderr or r.stdout or "Excel didn't save the file").strip().splitlines()[-1][:200])
    return out


def with_libreoffice(path, out_dir, timeout=180):
    from pathlib import Path
    soffice = find_libreoffice()
    os.makedirs(out_dir, exist_ok=True)
    profile = tempfile.mkdtemp(prefix="ds_lo_")   # own profile, so it works alongside an open LibreOffice
    try:
        ext = os.path.splitext(path)[1].lower().lstrip(".") or "xlsx"
        subprocess.run([soffice, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--norestore",
                        "--calc", "--convert-to", ext, "--outdir", out_dir, path],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=timeout, check=False,
                       env={**os.environ, "SAL_USE_VCLPLUGIN": "svp"})
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = os.path.join(out_dir, os.path.basename(path))
    if not (os.path.exists(out) and os.path.getsize(out)):
        raise RuntimeError("LibreOffice didn't produce a file")
    return out


def with_pycel(path, out_dir, cells):
    """Calculate `cells` ({sheet: [refs]}) with pycel; save a copy where every formula is a value.

    Returns (path, calculated, {sheet: cells left empty}, stopped_early)."""
    from openpyxl import load_workbook
    from pycel import ExcelCompiler
    logging.getLogger("pycel").setLevel(logging.CRITICAL)
    os.makedirs(out_dir, exist_ok=True)
    xl = ExcelCompiler(filename=path)
    wb = load_workbook(path)                    # formulas and formatting: we write values into it
    vals = load_workbook(path, data_only=True)  # saved results of the other formulas
    t0, done, empty, stopped = time.time(), 0, {}, False
    for sheet, refs in cells.items():
        for k, ref in enumerate(refs):
            if time.time() - t0 > PYCEL_TIME_LIMIT:
                stopped = True
                empty[sheet] = empty.get(sheet, 0) + len(refs) - k
                for later, rest in list(cells.items())[list(cells).index(sheet) + 1:]:
                    empty[later] = empty.get(later, 0) + len(rest)
                break
            try:
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    v = xl.evaluate(f"'{sheet}'!{ref}")
                if isinstance(v, (tuple, list)):   # a range result: keep the top-left value
                    while isinstance(v, (tuple, list)) and v:
                        v = v[0]
                wb[sheet][ref].value = v
                done += 1
            except Exception:
                wb[sheet][ref].value = None
                empty[sheet] = empty.get(sheet, 0) + 1
        if stopped:
            break
    # every other formula: keep its saved result as a plain value (openpyxl would drop it on save)
    pending = {(s, r) for s, refs in cells.items() for r in refs}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and (ws.title, c.coordinate) not in pending:
                    c.value = vals[ws.title][c.coordinate].value
    out = os.path.join(out_dir, os.path.basename(path))
    wb.save(out)
    return out, done, empty, stopped


def calculate(path, out_dir, missing, engines=None):
    """Try each engine in turn. `missing` is uncached() for this file.

    Returns {"path": calculated copy or None, "engine": name or None, "note": text for the user}."""
    n = sum(missing.values())
    tried = []
    for engine in engines or ENGINES:
        try:
            if engine == "excel" and find_excel():
                out = with_excel(path, os.path.join(out_dir, "recalc_excel"))
            elif engine == "libreoffice" and find_libreoffice():
                out = with_libreoffice(path, os.path.join(out_dir, "recalc_lo"))
            elif engine == "pycel" and find_pycel():
                size_mb = os.path.getsize(path) / 1024 / 1024
                if size_mb > PYCEL_MAX_MB:
                    tried.append(f"the built-in calculator only handles files up to {PYCEL_MAX_MB:g} MB (this one is {size_mb:.0f} MB)")
                    continue
                cells = uncached_cells(path, PYCEL_MAX_FORMULAS)
                if cells is None:
                    tried.append(f"the built-in calculator only handles up to {PYCEL_MAX_FORMULAS:,} such formulas "
                                 f"(this file has about {n:,})")
                    continue
                out, done, empty, stopped = with_pycel(path, os.path.join(out_dir, "recalc_pycel"), cells)
                total = sum(len(v) for v in cells.values())
                bits = []
                if not stopped and empty:
                    bits.append(f"{sum(empty.values()):,} use functions it doesn't support and are empty")
                if stopped:
                    bits.append(f"it stopped after {PYCEL_TIME_LIMIT:g} s; {sum(empty.values()):,} are empty")
                note = (f"{total:,} formula cell{'s' if total != 1 else ''} had no saved result; they were calculated with "
                        f"the built-in calculator" + (" (" + "; ".join(bits) + ")" if bits else "") + ".")
                log.info("%s: %s", os.path.basename(path), note)
                return {"path": out, "engine": "pycel", "note": note, "partial": bool(empty), "empty": empty}
            else:
                continue
            note = (f"{n:,} formula cell{'s' if n != 1 else ''} had no saved result (the file was saved by a program that "
                    f"doesn't calculate formulas); they were calculated with {NAMES[engine]}.")
            log.info("%s: %s", os.path.basename(path), note)
            return {"path": out, "engine": engine, "note": note, "partial": False, "empty": {}}
        except Exception as e:
            log.warning("%s couldn't calculate %s: %s", NAMES.get(engine, engine), os.path.basename(path), _short(e))
            tried.append(f"{NAMES.get(engine, engine)} failed ({_short(e)})")
    where = ", ".join(f"'{k}' ({v:,})" for k, v in list(missing.items())[:5])
    note = (f"{n:,} formula cell{'s' if n != 1 else ''} in sheet {where} ha{'ve' if n != 1 else 's'} no saved result, so "
            f"they're empty" + (": " + "; ".join(tried) if tried else "") +
            ". Open the file in Excel, save it and attach it again" +
            ("" if tried else ", or install LibreOffice on the server") + ".")
    return {"path": None, "engine": None, "note": note, "partial": False, "empty": dict(missing)}


def available():
    """Which engines this machine has, for the startup message."""
    found = [NAMES[e] for e in ENGINES if (e == "excel" and find_excel()) or (e == "libreoffice" and find_libreoffice())
             or (e == "pycel" and find_pycel())]
    return ", ".join(found) if found else "none"
