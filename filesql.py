"""
SQL over attached CSV, TSV and Excel files, using DuckDB.

At upload, each file (each sheet, for Excel) is loaded into a small DuckDB database stored next
to the upload (tables.duckdb), with column types detected from the data and column names
normalised to snake_case. At query time, the chat's files are attached read-only to a fresh
in-memory DuckDB, exposed as views with chat-level table names, and the session is locked down:

  1. Exactly one statement, and it must be a SELECT (or EXPLAIN of one).
  2. enable_external_access=false + lock_configuration=true: no reading or writing files, no
     ATTACH, no extensions, and the settings can't be changed back.
  3. The data files are attached READ_ONLY, and the database itself is a throwaway in-memory one.
  4. A timeout (interrupt) and a row cap, like PostgreSQL queries.

DuckDB's SQL is close to PostgreSQL's (::casts, ILIKE, date_trunc, string_agg, window functions).

Several tables in one sheet or CSV are split into separate tables (see _split): blocks separated by
blank rows (stacked) or blank columns (side by side). A one-cell title row names the table below it.
A block that doesn't start with a header row, or repeats the header above it, continues that table.
Ordinary one-table files never go through the splitter, so they load exactly as before.
"""
from array import array
import csv
import datetime
import logging
import mmap
import os
import re
import threading
import time

try:
    import duckdb
except ImportError:  # optional dependency
    duckdb = None

log = logging.getLogger("datasquint.files")

TABLE_EXTS = {".csv", ".tsv", ".xlsx", ".xlsm"}
EXCEL_EXTS = {".xlsx", ".xlsm"}
HEAD_BYTES = 64 * 1024      # how much of each file the text preview reads
DB_FILE = "tables.duckdb"
SOURCE = "files"            # what the model writes in "-- db: files"
SAMPLE_ROWS = 3
PREVIEW_LINES = 15
MEMORY_LIMIT = os.getenv("FILES_SQL_MEMORY", "1GB")
# "calamine" (default, ~10x faster, ~1 GB RAM per million-row sheet) or "openpyxl" (slow, low memory)
EXCEL_READER = os.getenv("EXCEL_READER", "calamine").strip().lower()
# Split sheets/CSVs holding several tables into one table each (1 = on, the default)
SPLIT_TABLES = os.getenv("SPLIT_TABLES", "1").strip().lower() not in ("0", "false", "no", "off")
MAX_SPLIT_TABLES = int(os.getenv("MAX_SPLIT_TABLES", 30))   # more blocks than this: keep it as one table


# FILE_TABLES=0 turns SQL over attached CSV/TSV/Excel off: those files are then read as text,
# under the normal MAX_UPLOAD_MB limit, like any other document.
ENABLED = os.getenv("FILE_TABLES", "1").strip().lower() not in ("0", "false", "no", "off")


def installed():
    return duckdb is not None


def available():
    """True when attached CSV/TSV/Excel become SQL tables: DuckDB installed and FILE_TABLES on."""
    return installed() and ENABLED


def status():
    if not installed():
        return "off (DuckDB isn't installed: pip install duckdb)"
    return "on" if ENABLED else "off (FILE_TABLES=0)"


def is_tabular(name):
    return os.path.splitext(name)[1].lower() in TABLE_EXTS


def _ident(name):
    n = re.sub(r"[^a-z0-9_]+", "_", name.lower()).strip("_") or "table"
    if n[0].isdigit():
        n = "t_" + n
    return n[:48]


def _q(name):
    return '"' + name.replace('"', '""') + '"'


# ---------------------------------------------------------------- upload time
def _encoding(path):
    """The encoding a CSV decodes with: UTF-8, else Windows-1252 (common from Excel), else Latin-1."""
    import codecs
    for enc in ("utf-8-sig", "cp1252"):
        dec = codecs.getincrementaldecoder(enc)()
        try:
            with open(path, "rb") as fh:
                while chunk := fh.read(1 << 20):
                    dec.decode(chunk)
                dec.decode(b"", final=True)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def _utf8_copy(path, folder, idx):
    """Copy a CSV to UTF-8 (Excel often saves Windows-1252)."""
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    out = os.path.join(folder, f"source_{idx}_utf8.csv")
    with open(out, "w", encoding="utf-8", newline="") as fh:
        fh.write(text)
    return out


EXCEL_ERRORS = {"#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!", "#GETTING_DATA"}


def _xcell(v):
    """An Excel cell value as CSV text that DuckDB will type correctly."""
    # Both Excel readers (python-calamine and openpyxl) give the same CSV text: error cells are empty
    # (calamine already does that), and a date-time at exactly midnight is a date (calamine returns dates).
    if v is None or (isinstance(v, str) and v in EXCEL_ERRORS):
        return ""
    if isinstance(v, datetime.datetime):
        if not (v.hour or v.minute or v.second or v.microsecond):
            return v.date().isoformat()
        return v.isoformat(sep=" ")
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, datetime.timedelta):
        return str(v)
    if isinstance(v, float) and v.is_integer() and abs(v) < 2 ** 53:
        return int(v)  # Excel stores every number as a float; keep whole numbers as integers
    return v


def _fatal(e):
    return isinstance(e, (KeyboardInterrupt, SystemExit))


def _calamine_book(path):
    """A python-calamine workbook, or None (not installed, turned off, or it can't open the file)."""
    if EXCEL_READER == "openpyxl":
        return None
    try:
        import python_calamine
        return python_calamine.CalamineWorkbook.from_path(path)
    except BaseException as e:   # ImportError, or a Rust panic (pyo3 PanicException is a BaseException)
        if _fatal(e):
            raise
        return None


class _Openpyxl:
    """openpyxl, opened on first use (it's slower, so only when calamine can't do a sheet)."""
    def __init__(self, path):
        self.path, self.wb, self.error = path, None, None

    def book(self):
        if self.wb is None and self.error is None:
            try:
                from openpyxl import load_workbook
                self.wb = load_workbook(self.path, read_only=True, data_only=True)
            except Exception as e:
                self.error = e
        if self.error is not None:
            raise self.error
        return self.wb

    def close(self):
        if self.wb is not None:
            self.wb.close()


def _write_sheet(rows, dest):
    """Write a sheet's rows to CSV, keeping blank rows as empty lines (DuckDB skips them; the splitter needs
    the gaps). Returns (data rows written, rowmap, width)."""
    n = width = 0
    rowmap = array("I")
    with open(dest, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        blank = 0
        for r, row in enumerate(rows, start=1):
            if not any(v is not None and str(v).strip() != "" for v in row):
                blank += 1
                continue
            if n and blank:
                gap = min(blank, 2)
                fh.write("\n" * gap)
                rowmap.extend([r - 1] * gap)
            blank = 0
            w.writerow([_xcell(v) for v in row])
            width = max(width, len(row))
            rowmap.append(r)
            n += 1
    return n, rowmap, width


def _xlsx_sheets(path, folder, warnings=None):
    """Write each non-empty sheet to a CSV.

    Each sheet is read with python-calamine (fast) and, if that fails for that sheet, again with openpyxl,
    so one awkward sheet never costs the others. Sheets neither can read are reported in `warnings`.

    Returns [(sheet_name, csv_path, coords)], coords = (rowmap, col_offset, width): rowmap[k] is the Excel row
    number (1-based) of CSV record k, and column j of the CSV is Excel column j + col_offset + 1. They let
    a generated script find each table in the workbook itself.
    """
    warnings = warnings if warnings is not None else []
    cal, opx = _calamine_book(path), _Openpyxl(path)
    try:
        names = cal.sheet_names if cal else opx.book().sheetnames
    except Exception as e:
        raise ValueError(_excel_error(e)) from None
    out = []
    try:
        for i, name in enumerate(names):
            dest = os.path.join(folder, f"sheet_{i}.csv")
            done, last = False, None
            if cal:
                try:
                    sheet = cal.get_sheet_by_name(name)
                    if not sheet.height:   # no cells (calamine panics on iter_rows() for these)
                        continue
                    n, rowmap, width = _write_sheet(sheet.iter_rows(), dest)
                    off = sheet.start[1] if sheet.start else 0   # calamine skips empty leading columns
                    done = True
                except BaseException as e:
                    if _fatal(e):
                        raise
                    last = e
                    log.info("python-calamine couldn't read sheet %r of %s (%s); trying openpyxl", name,
                             os.path.basename(path), _short(e))
            if not done:
                try:
                    n, rowmap, width = _write_sheet(opx.book()[name].iter_rows(values_only=True), dest)
                    off, done = 0, True
                except Exception as e:
                    last = e
            if not done:
                warnings.append(f"Sheet '{name}' couldn't be read and was skipped ({_short(last)}).")
                if os.path.exists(dest):
                    os.remove(dest)
                continue
            if n:
                out.append((name, dest, (rowmap, off, width)))
            else:
                os.remove(dest)
    finally:
        opx.close()
    return out


def _short(e):
    text = str(e).strip().splitlines()[0] if str(e).strip() else type(e).__name__
    return text[:200]


def _excel_error(e):
    """A plain explanation for a workbook that can't be opened at all."""
    text = str(e)
    if "zip" in text.lower() or "BadZipFile" in type(e).__name__:
        return ("This isn't a readable .xlsx file. It may be password-protected (remove the password in Excel), "
                "or an older .xls file renamed to .xlsx (save it as .xlsx or CSV in Excel).")
    return f"Couldn't read this Excel file: {_short(e)}"


# ---------------------------------------------------------------- several tables in one sheet
_BLANK_THEN_DATA = re.compile(rb"\n[ \t,;|\"']*\r?\n(?=[ \t,;|\"']*[^ \t,;|\"'\r\n])")
_AUTO_COL = re.compile(r"^column\d+$")
_NUMISH = re.compile(r"^[-+(]?[$€£¥]?\s*[\d.,]+\s*[%)]?$|^\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}([ T]\d{1,2}:\d{2}(:\d{2})?)?$|^\d{1,2}:\d{2}(:\d{2})?$")


def _has_gaps(src):
    """A blank line with data after it?"""
    try:
        with open(src, "rb") as fh, mmap.mmap(fh.fileno(), 0, access=mmap.ACCESS_READ) as mm:
            return _BLANK_THEN_DATA.search(mm) is not None
    except (ValueError, OSError):   # empty file
        return False


def _maybe_several(src, cols):
    """Cheap check: could this source hold more than one table (or a title row)?"""
    # an empty header cell (a gap column or a title row), or a blank line with data after it
    return any(_AUTO_COL.match(c) for c, _ in cols) or _has_gaps(src)


def _header_like(cells):
    vals = [c.strip() for c in cells if c.strip()]
    return bool(vals) and not any(_NUMISH.match(v) for v in vals)


def _shape(cells):
    """Rough pattern of a row ("SN003" -> "aa9"), to tell more data from a new header."""
    return [re.sub(r"(9)9+|(a)a+", r"\1\2", re.sub(r"\d", "9", re.sub(r"[^\W\d_]", "a", c.strip()))) for c in cells]


def _csv_rows(path, delim):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        yield from csv.reader(fh, delimiter=delim)


def _sniff_delim(path, delim):
    if delim:
        return delim
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as fh:
        head = fh.read(64 * 1024)
    try:
        return csv.Sniffer().sniff(head, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _split(src, delim, folder, idx):
    """Find the tables in a CSV: blocks between blank rows, split again at blank columns.

    Returns [{"path", "title", "header"}] for two or more tables (or one table that lost a title or
    repeated header rows), or None when it's simply one table and the normal load is right.
    Two streaming passes, so memory stays small even for big files.
    """
    try:
        try:
            return _split_inner(src, delim, folder, idx)
        except UnicodeDecodeError:
            return _split_inner(_utf8_copy(src, folder, idx), delim, folder, idx)
    except Exception:   # never let the splitter break an upload: fall back to the normal load
        log.warning("Couldn't look for several tables in %s; loading it as one table", os.path.basename(src), exc_info=True)
        return None


def _no_blank_lines(src, folder, idx):
    """Copy of a sheet CSV without its empty lines (how sheets were loaded before splitting existed)."""
    out = os.path.join(folder, f"source_{idx}_rows.csv")
    with open(src, encoding="utf-8", newline="") as fi, open(out, "w", encoding="utf-8", newline="") as fo:
        for line in fi:
            if line.strip():
                fo.write(line)
    return out


def _split_inner(src, delim, folder, idx):
    delim = _sniff_delim(src, delim)
    # pass 1: segments of non-blank rows, the columns each uses, and their first few rows
    segs, cur = [], None
    for i, row in enumerate(_csv_rows(src, delim)):
        used = [j for j, c in enumerate(row) if c.strip()]
        if not used:
            cur = None
            continue
        if cur is None:
            cur = {"start": i, "end": i, "cols": set(), "head": []}
            segs.append(cur)
            if len(segs) > 50 * MAX_SPLIT_TABLES:
                return None
        cur["end"] = i
        cur["cols"].update(used)
        if len(cur["head"]) < 5:
            cur["head"].append((i, row))

    # blocks: each segment split at fully empty columns, left to right
    blocks = []
    for sg in segs:
        cols = sorted(sg["cols"])
        groups, lo = [], cols[0]
        for a, b in zip(cols, cols[1:]):
            if b - a > 1:
                groups.append((lo, a))
                lo = b
        groups.append((lo, cols[-1]))
        for lo, hi in groups:
            first = [(i, [c.strip() for c in (row[lo:hi + 1] + [""] * (hi + 1 - len(row[lo:hi + 1])))])
                     for i, row in sg["head"] if any(c.strip() for c in row[lo:hi + 1])]
            blocks.append({"seg": sg, "lo": lo, "hi": hi, "first": first})

    # plan: decide which blocks start a table, which continue one, and title/header rows to skip
    tables, open_at, pending_title, changed = [], {}, None, False
    for bk in blocks:
        first, lo, hi = bk["first"], bk["lo"], bk["hi"]
        skip = set()
        n_first = sum(1 for c in first[0][1] if c) if first else 0
        # a block that's only a one-cell line: a title (or note) for the next table
        prev = open_at.get(lo)
        seg_rows = bk["seg"]["end"] - bk["seg"]["start"] + 1
        # a short block of one-cell lines (a title, "Generated on …" notes): a title for the next table,
        # unless it looks like more rows of a one-column table above
        if (first and seg_rows <= 3 and len(first) == seg_rows and all(sum(1 for c in r if c) == 1 for _, r in first)
                and not (prev is not None and prev["data"] is not None and hi <= prev["hi"]
                         and _shape(first[0][1]) == _shape(prev["data"][:len(first[0][1])]))):
            pending_title, changed = first[0][1][[bool(c) for c in first[0][1]].index(True)], True
            continue
        title = None
        if first and len(first) >= 2 and n_first == 1 and sum(1 for c in first[1][1] if c) >= 2:
            title = next(c for c in first[0][1] if c)
            skip.add(first[0][0])
            first = first[1:]
            changed = True
        header = first[0][1] if first else None
        looks_like_data = header is not None and (
            not _header_like(header)
            or (prev is not None and prev["data"] is not None
                and _shape(header) == _shape(prev["data"][:len(header)])))
        if (prev is not None and title is None and pending_title is None and header is not None
                and hi <= prev["hi"] and (looks_like_data or header == prev["header"][:len(header)])):
            if header == prev["header"][:len(header)]:
                skip.add(first[0][0])   # the header repeated (e.g. page breaks)
                changed = True
            prev["parts"].append((bk, skip))
            continue
        t = {"lo": lo, "hi": hi, "header": header or [], "title": title or pending_title,
             "data": first[1][1] if len(first) > 1 else None,
             "header_row": _header_like(header or []), "parts": [(bk, skip)]}
        pending_title = None
        tables.append(t)
        open_at[lo] = t
        if len(tables) > MAX_SPLIT_TABLES:
            return None
    if not tables or (len(tables) == 1 and not changed):
        return None

    # pass 2: write each table's rows to its own CSV
    owner = {}   # id(segment) -> [(lo, hi, table, skip)]
    for t in tables:
        for bk, skip in t["parts"]:
            owner.setdefault(id(bk["seg"]), []).append((bk["lo"], bk["hi"], t, skip))
    files, writers = [], {}
    try:
        for k, t in enumerate(tables):
            path = os.path.join(folder, f"block_{idx}_{k}.csv")
            fh = open(path, "w", encoding="utf-8", newline="")
            files.append(fh)
            writers[id(t)] = csv.writer(fh)
            t["path"], t["rows"] = path, 0
        si = 0
        for i, row in enumerate(_csv_rows(src, delim)):
            while si < len(segs) and segs[si]["end"] < i:
                si += 1
            if si >= len(segs) or segs[si]["start"] > i:
                continue
            for lo, hi, t, skip in owner.get(id(segs[si]), ()):
                if i in skip:
                    continue
                cells = row[lo:hi + 1]
                if not any(c.strip() for c in cells):
                    continue
                width = t["hi"] - t["lo"] + 1
                writers[id(t)].writerow((cells + [""] * width)[:width])
                t["rows"] += 1
    finally:
        for fh in files:
            fh.close()
    out = []
    for t in tables:
        if t["rows"] < (2 if t["header_row"] else 1):   # a header with no data under it
            os.remove(t["path"])
            continue
        out.append({"path": t["path"], "title": t["title"], "header": t["header_row"],
                    "first": next((c for c in t["header"] if c), ""), "delim": delim,
                    "parts": [[bk["seg"]["start"], bk["seg"]["end"], bk["lo"], bk["hi"], sorted(skip)]
                              for bk, skip in t["parts"]]})
    if not out or (len(out) == 1 and not out[0]["title"] and not changed):
        for t in out:
            os.remove(t["path"])
        return None
    return out


def _load_spec(sheet, coords, encoding, part, delim):
    """How a generated script can rebuild this table from the original file (see scriptgen.py)."""
    if sheet:
        rowmap, col_offset, width = coords
        spec = {"kind": "sheet", "sheet": sheet, "col_offset": col_offset, "width": width}
        if part:   # CSV record numbers -> Excel rows, CSV columns -> Excel columns (1-based)
            spec["parts"] = [[rowmap[a], rowmap[b], lo + col_offset + 1, hi + col_offset + 1, [rowmap[k] for k in skip]]
                             for a, b, lo, hi, skip in part["parts"]]
    else:
        spec = {"kind": "csv", "encoding": encoding, "delim": delim}
        if part:
            spec.update(delim=part["delim"], parts=part["parts"])
    if part:
        spec["header"] = part["header"]
    return spec


def _label(part, k):
    """Short name for a table found inside a sheet: its title, else its first column heading."""
    for text in (part.get("title"), part.get("first") if part.get("header") else None):
        if text and _ident(text) != "table":
            words, label = _ident(text).split("_"), ""
            for w in words:
                if label and len(label) + 1 + len(w) > 24:
                    break
                label = f"{label}_{w}" if label else w[:24]
            return label
    return f"table_{k + 1}"


BASE_OPTS = "normalize_names=true, sample_size=-1, null_padding=true"
OUR_CSV = "delim=',', quote='\"', escape='\"'"   # CSVs Datasquint wrote itself (sheets, split tables)
TEXT_NOTE = ("Some columns had mixed or unreadable values, so all columns were loaded as text; "
             "cast them in SQL (e.g. qty::INTEGER, day::DATE).")
SKIP_NOTE = "Some malformed lines couldn't be read and were left out."


def _attempts(dialect, extra=""):
    """read_csv settings to try in turn, each more forgiving: (options, note for the model or None)."""
    first = ", ".join(x for x in (dialect, extra) if x)
    loose = ", ".join(x for x in (dialect or "quote='\"'", extra, "strict_mode=false", "max_line_size=200000000") if x)
    return [(first, None), (loose, None),
            (loose + ", all_varchar=true", TEXT_NOTE),
            (loose + ", all_varchar=true, ignore_errors=true", TEXT_NOTE + " " + SKIP_NOTE)]


def _create(con, table, path, attempts):
    """CREATE TABLE from a CSV with the first settings that work. Returns (options used, note)."""
    last = None
    for k, (opts, note) in enumerate(attempts):
        try:
            con.execute(f"DROP TABLE IF EXISTS {table}")
            con.execute(f"CREATE TABLE {table} AS SELECT * FROM read_csv(?, {', '.join(x for x in (BASE_OPTS, opts) if x)})", [path])
            if k:
                log.info("%s loaded with fallback settings (%s) after: %s", os.path.basename(path), opts, _short(last))
            return opts, note
        except duckdb.Error as e:   # e.g. an option this DuckDB version doesn't know: just try the next
            last = e
    con.execute(f"DROP TABLE IF EXISTS {table}")
    raise ValueError(f"DuckDB couldn't read it as a table: {_short(last)}")


def prepare(path, name, folder, warnings=None):
    """Load a tabular upload into folder/tables.duckdb.

    Returns [{"table", "sheet", "columns": [[name, type]], "rows", "load", ...}]. Raises ValueError with a
    plain reason when nothing in the file can be loaded. Problems that didn't stop it (a sheet skipped,
    columns loaded as text) are appended to `warnings`.
    """
    warnings = warnings if warnings is not None else []
    ext = os.path.splitext(name)[1].lower()
    if ext in EXCEL_EXTS:
        try:
            sources = [(sheet, p, None, coords) for sheet, p, coords in _xlsx_sheets(path, folder, warnings)]
        except ValueError:
            raise
        except BaseException as e:  # don't let a reader panic take down the request
            if _fatal(e):
                raise
            raise ValueError(_excel_error(e)) from None
        if not sources:
            raise ValueError("No sheet in this workbook has any data." if not warnings else " ".join(warnings))
        encoding = None
    else:
        encoding = _encoding(path)
        if not encoding.startswith("utf-8"):   # e.g. Windows-1252 from Excel: load a UTF-8 copy
            src = _utf8_copy(path, folder, 0)
        else:
            src = path
        sources = [("", src, "\t" if ext == ".tsv" else None, None)]
    dbpath = os.path.join(folder, DB_FILE)
    con = duckdb.connect(dbpath)
    out = []
    try:
        for i, (sheet, src, delim, coords) in enumerate(sources):
            table = f"t{i}"
            where = f"sheet '{sheet}'" if sheet else "the file"
            dialect = OUR_CSV if sheet else ("delim='\\t'" if delim else "")
            try:
                opts, note = _create(con, table, src, _attempts(dialect))
                loaded = True
            except ValueError as e:
                loaded, opts, note, err = False, "", None, e   # the splitter may still manage (e.g. a title row)
            cols = [[c[0], c[1]] for c in con.execute(f"DESCRIBE {table}").fetchall()] if loaded else []
            parts = (_split(src, delim, folder, i)
                     if SPLIT_TABLES and (not loaded or _maybe_several(src, cols)) else None)
            if loaded and not parts and sheet and _has_gaps(src):
                # one table after all: load it without the blank rows kept for the splitter (in a one-column
                # sheet DuckDB would read them as empty rows)
                try:
                    opts, note = _create(con, table, _no_blank_lines(src, folder, i), _attempts(dialect))
                    cols = [[c[0], c[1]] for c in con.execute(f"DESCRIBE {table}").fetchall()]
                except ValueError:
                    opts, note = _create(con, table, src, _attempts(dialect))
            if not loaded and not parts:
                if sheet:   # Excel: one more try without the blank rows kept for the splitter
                    try:
                        opts, note = _create(con, table, _no_blank_lines(src, folder, i), _attempts(dialect))
                        cols = [[c[0], c[1]] for c in con.execute(f"DESCRIBE {table}").fetchall()]
                    except ValueError as e:
                        warnings.append(f"Sheet '{sheet}' couldn't be loaded as a table ({e}).")
                        continue
                else:
                    raise err
            if parts:
                con.execute(f"DROP TABLE IF EXISTS {table}")
                labels = set()
                for k, part in enumerate(parts):
                    sub = f"{table}_{k}"
                    hdr = "header=true" if part["header"] else ""
                    try:
                        popts, pnote = _create(con, sub, part["path"], _attempts(OUR_CSV, hdr))
                    except ValueError as e:
                        warnings.append(f"A table in {where} couldn't be loaded ({e}).")
                        continue
                    label, n = _label(part, k), 2
                    while label in labels:
                        label, n = f"{_label(part, k)}_{n}", n + 1
                    labels.add(label)
                    spec = _load_spec(sheet, coords, encoding, part, delim)
                    spec["opts"] = popts
                    out.append({"table": sub, "sheet": sheet, "block": label, "title": part["title"] or "",
                                "load": spec, **({"note": pnote} if pnote else {}),
                                "columns": [[c[0], c[1]] for c in con.execute(f"DESCRIBE {sub}").fetchall()],
                                "rows": int(con.execute(f"SELECT count(*) FROM {sub}").fetchone()[0]),
                                "src": os.path.basename(part["path"])})
                    if pnote:
                        warnings.append(f"{label}: {pnote}")
                continue
            rows = con.execute(f"SELECT count(*) FROM {table}").fetchone()[0]
            if cols:
                spec = _load_spec(sheet, coords, encoding, None, delim)
                spec["opts"] = opts
                out.append({"table": table, "sheet": sheet, "columns": cols, "rows": int(rows), "load": spec,
                            **({"note": note} if note else {}), "src": os.path.basename(src)})
                if note:
                    warnings.append(f"{where[0].upper() + where[1:]}: {note}")
    finally:
        con.close()
    if not out:
        if os.path.exists(dbpath):
            os.remove(dbpath)
        raise ValueError(" ".join(warnings) or "No data found to load as a table.")
    return out


def off_reason():
    """Why CSV/Excel files aren't loaded as tables on this server, or None if they are."""
    if not installed():
        return "DuckDB isn't installed on the server (pip install duckdb), so CSV and Excel files are read as text."
    if not ENABLED:
        return "Loading CSV and Excel files as tables is turned off on this server (FILE_TABLES=0)."
    return None


# ---------------------------------------------------------------- chat level
def chat_tables(file_ids, load_meta, upload_dir):
    """Tables for the chat's files, in order, with unique chat-level names.

    Returns [{"name", "fid", "file", "sheet", "table", "columns", "rows", "db"}].
    The same ordered ids always give the same names, so the prompt and the queries agree.
    """
    out, used = [], set()
    for fid in file_ids:
        meta = load_meta(fid)
        if not meta or not meta.get("tables"):
            continue
        db = os.path.join(upload_dir, fid, DB_FILE)
        if not os.path.exists(db):
            continue
        base = _ident(os.path.splitext(meta["name"])[0])
        sheets = {t.get("sheet") for t in meta["tables"]}
        for t in meta["tables"]:
            name = base if not t.get("sheet") or len(sheets) == 1 else _ident(f"{base}_{t['sheet']}")
            if t.get("block"):   # one of several tables found in the same sheet/file
                name = _ident(f"{name}_{t['block']}")
            unique, n = name, 2
            while unique in used:
                unique, n = f"{name}_{n}", n + 1
            used.add(unique)
            out.append(dict(t, name=unique, fid=fid, file=meta["name"], db=db))
    return out


def _connect(tables):
    con = duckdb.connect(":memory:")
    con.execute(f"SET memory_limit='{MEMORY_LIMIT}'")
    con.execute("SET threads=2")
    aliases = {}
    for t in tables:
        if t["db"] not in aliases:
            alias = f"f{len(aliases)}"
            con.execute(f"ATTACH '{t['db'].replace(chr(39), chr(39) * 2)}' AS {alias} (READ_ONLY)")
            aliases[t["db"]] = alias
        con.execute(f"CREATE VIEW {_q(t['name'])} AS SELECT * FROM {aliases[t['db']]}.{t['table']}")
    con.execute("SET enable_external_access=false")
    con.execute("SET lock_configuration=true")
    return con


def _cell(v, full=False):
    cap = None if full else 2000
    if v is None or isinstance(v, (bool, int, str)):
        return v[:cap] if isinstance(v, str) else v
    if isinstance(v, float):
        return v if v == v and abs(v) != float("inf") else str(v)
    return str(v)[:cap]


def run_query(sql, tables, max_rows, timeout_ms, full=False):
    """Run one read-only SELECT over the chat's file tables."""
    if not tables:
        raise ValueError("This chat has no attached CSV or Excel files to query.")
    t0 = time.time()
    con = _connect(tables)
    timer = threading.Timer(timeout_ms / 1000, con.interrupt)
    try:
        stmts = con.extract_statements(sql)
        if len(stmts) != 1:
            raise ValueError("Run one statement at a time.")
        if stmts[0].type not in (duckdb.StatementType.SELECT, duckdb.StatementType.EXPLAIN):
            raise ValueError("Only read-only queries can be run on attached files (SELECT, WITH, EXPLAIN).")
        timer.start()
        try:
            cur = con.execute(sql)
            columns = [d[0] for d in cur.description] if cur.description else []
            rows = cur.fetchmany(max_rows + 1) if columns else []
        except duckdb.InterruptException:
            raise ValueError(f"The query took longer than {timeout_ms // 1000}s and was stopped.")
    finally:
        timer.cancel()
        con.close()
    truncated = len(rows) > max_rows
    return {"db_id": SOURCE, "db_name": SOURCE, "columns": columns,
            "rows": [[_cell(v, full) for v in r] for r in rows[:max_rows]], "row_count": min(len(rows), max_rows),
            "truncated": truncated, "ms": int((time.time() - t0) * 1000)}


# ---------------------------------------------------------------- prompt
def _md(v):
    s = "NULL" if v is None else str(v)
    return s.replace("|", "\\|").replace("\n", " ")[:60]


def prompt_section(tables, max_chars=20000):
    """Schema of the chat's file tables, with a few sample rows each."""
    lines, samples = [], {}
    try:
        con = _connect(tables)
        try:
            for t in tables:
                samples[t["name"]] = con.execute(f"SELECT * FROM {_q(t['name'])} LIMIT {SAMPLE_ROWS}").fetchall()
        finally:
            con.close()
    except Exception:
        pass
    for t in tables:
        src = t["file"] + (f", sheet '{t['sheet']}'" if t.get("sheet") else "")
        if t.get("block"):
            src += ", one of several tables in it" + (f", titled '{t['title']}'" if t.get("title") else "")
        block = [f"{t['name']}  (from {src}; {t['rows']:,} rows)"] + ([f"  note: {t['note']}"] if t.get("note") else [])
        block += [f"  {c} {typ}" for c, typ in t["columns"]]
        rows = samples.get(t["name"])
        if rows:
            block.append("  sample rows:")
            block.append("  | " + " | ".join(c for c, _ in t["columns"]) + " |")
            block += ["  | " + " | ".join(_md(v) for v in r) + " |" for r in rows]
        text = "\n".join(block)
        if sum(len(x) for x in lines) + len(text) > max_chars:
            lines.append(f"... {len(tables) - tables.index(t)} more table(s) not shown")
            break
        lines.append(text)
    return ("Files attached to this chat, loaded as tables you can query with SQL. This source uses DuckDB, "
            "whose SQL is close to PostgreSQL (::casts, ILIKE, date_trunc, string_agg and window functions work). "
            "Column names were normalised to lower_snake_case. Prefer SQL over reading the file text for counts, "
            "totals, averages, rankings and filters; the file text you see is only a preview.\n\n"
            + "\n\n".join(lines))


def _read_head(path, max_lines):
    """First lines of a text file, without reading the whole (possibly huge) file."""
    with open(path, "rb") as fh:
        raw = fh.read(HEAD_BYTES)
    cut = len(raw) == HEAD_BYTES
    text = None
    for trim in (0, 1, 2, 3) if cut else (0,):  # a cut can split a UTF-8 character
        try:
            text = raw[:len(raw) - trim].decode("utf-8-sig")
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        try:
            text = raw.decode("cp1252")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
    lines = text.splitlines()
    if cut:
        lines = lines[:-1]  # the last line may be incomplete
    return "\n".join(lines[:max_lines])


def head_text(folder, tables, max_lines=200):
    """Text kept for a tabular upload: the first lines of each table's source, not the whole file."""
    parts = []
    for t in tables:
        head = _read_head(os.path.join(folder, t["src"]), max_lines)
        label = " · ".join(x for x in (f"Sheet: {t['sheet']}" if t.get("sheet") else "",
                                        f"Table: {t.get('title') or t['block']}" if t.get("block") else "") if x)
        parts.append(f"[{label}]\n{head}" if label else head)
    return "\n".join(parts)


def cleanup(folder):
    """Delete intermediate CSVs (Excel sheets, UTF-8 copies) once the tables and preview exist."""
    for f in os.listdir(folder):
        if f.startswith(("sheet_", "source_", "block_")) and f.endswith(".csv"):
            try:
                os.remove(os.path.join(folder, f))
            except OSError:
                pass


def preview(text, tables):
    """Short preview of a tabular file for the message body, instead of the full text."""
    head = "\n".join(text.splitlines()[:PREVIEW_LINES])
    names = ", ".join(f"{t['name']} ({t['rows']:,} rows)" for t in tables)
    return f"[Loaded as SQL table(s): {names}. Only the first {PREVIEW_LINES} lines are shown here; query the table for the rest.]\n{head}"
